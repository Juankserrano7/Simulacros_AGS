import re
from functools import reduce
import unicodedata

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_") or "sim"


def _normalize_name(name: str) -> str:
    if not isinstance(name, str):
        return ""
    # Quita acentos, trim y colapsa espacios
    norm = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    norm = re.sub(r"[^A-Z0-9 ]+", " ", norm.upper())
    norm = re.sub(r"\s+", " ", norm).strip()
    return norm


def render(simulacros, materias):
    st.markdown("<h2 class='section-header'>🏆 Top Global</h2>", unsafe_allow_html=True)

    # Unificar datos de todos los simulacros
    frames = []
    grado_cols = []
    grados_list = []
    for sim in simulacros:
        base = sim["df"][["ESTUDIANTE"] + materias + ["PROMEDIO PONDERADO"]].copy()
        base["ESTUDIANTE"] = base["ESTUDIANTE"].apply(_normalize_name)
        slug = _slug(sim["nombre"])
        rename_map = {
            "PROMEDIO PONDERADO": f"{slug}_PROM",
            **{mat: f"{mat}_{slug}" for mat in materias},
        }
        if "GRADO" in sim["df"].columns:
            grado_col = f"GRADO_{slug}"
            base[grado_col] = sim["df"]["GRADO"]
            grado_cols.append(grado_col)
            grados_list.extend(sim["df"]["GRADO"].dropna().astype(str).tolist())
        base = base.rename(columns=rename_map)
        frames.append(base)

    datos_unificados = reduce(lambda left, right: left.merge(right, on="ESTUDIANTE", how="outer"), frames)

    prom_cols = [c for c in datos_unificados.columns if c.endswith("_PROM")]
    datos_unificados["PROMEDIO_PONDERADO_GENERAL"] = datos_unificados[prom_cols].mean(axis=1, skipna=True)
    for mat in materias:
        cols_mat = [c for c in datos_unificados.columns if c.startswith(mat + "_")]
        if cols_mat:
            datos_unificados[f"{mat}_PROMEDIO_GENERAL"] = datos_unificados[cols_mat].mean(axis=1, skipna=True)

    datos_unificados["SIMULACROS_PRESENTADOS"] = datos_unificados[prom_cols].notna().sum(axis=1)

    # Consolidar grado si hay múltiples columnas
    grado_candidates = [c for c in datos_unificados.columns if c.startswith("GRADO_")]
    if grado_candidates:
        datos_unificados["GRADO"] = datos_unificados[grado_candidates].bfill(axis=1).iloc[:, 0]
        datos_unificados.drop(columns=grado_candidates, inplace=True)

    # --- Top Global visual estilo previo ---
    top_global = (
        datos_unificados[["ESTUDIANTE", "PROMEDIO_PONDERADO_GENERAL"]]
        .dropna()
        .sort_values("PROMEDIO_PONDERADO_GENERAL", ascending=False)
        .head(30)
        .reset_index(drop=True)
    )
    if top_global.empty:
        st.warning("⚠️ No se encontraron datos válidos para el ranking global.")
    else:
        st.markdown("<div>", unsafe_allow_html=True)
        for idx in range(len(top_global)):
            estudiante = top_global.iloc[idx]
            nombre_completo = str(estudiante["ESTUDIANTE"]).strip()
            nombre_partes = nombre_completo.split()
            nombre_corto = " ".join(nombre_partes[:2]) if len(nombre_partes) > 1 else nombre_completo

            posicion_display = f"{idx + 1}"
            background_color = "linear-gradient(135deg, #667eea, #764ba2)"
            font_size = "1.2rem"

            if idx == 0:
                posicion_display = "🥇"
                background_color = "linear-gradient(135deg, #FFD700, #FFA500)"
                font_size = "1.5rem"
            elif idx == 1:
                posicion_display = "🥈"
                background_color = "linear-gradient(135deg, #C0C0C0, #E8E8E8)"
                font_size = "1.4rem"
            elif idx == 2:
                posicion_display = "🥉"
                background_color = "linear-gradient(135deg, #CD7F32, #E8C39E)"
                font_size = "1.3rem"

            st.markdown(
                f"""
            <div style="display: flex; align-items: center; padding: 1rem; margin-bottom: 0.5rem; background: linear-gradient(90deg, #f8f9fa 0%, #e9ecef 100%); border-radius: 10px; transition: transform 0.2s;">
                <div style="width: 50px; height: 50px; background: {background_color}; border-radius: 10px; display: flex; align-items: center; justify-content: center; font-weight: 800; font-size: {font_size}; color: white; margin-right: 1rem;">
                    {posicion_display}
                </div>
                <div style="flex: 1;">
                    <div style="font-weight: 600; font-size: 1rem; color: #2c3e50; margin-bottom: 0.3rem;" title="{nombre_completo}">{nombre_corto}</div>
                </div>
                <div style="font-size: 1.5rem; font-weight: 800; color: #667eea; margin-right: 1rem;">
                    {estudiante['PROMEDIO_PONDERADO_GENERAL']:.1f}
                </div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        st.markdown("</div>", unsafe_allow_html=True)

    st.markdown("---")
    st.markdown("<h2 class='section-header'>📋 Tabla Completa - Ranking Global por Estudiante</h2>", unsafe_allow_html=True)
    st.markdown("### 🎯 Controles de Visualización")

    col1, col2, col3, col4 = st.columns(4)
    metricas_materias = [f"{m}_PROMEDIO_GENERAL" for m in materias if f"{m}_PROMEDIO_GENERAL" in datos_unificados.columns]
    with col1:
        metrica_ordenar = st.selectbox(
            "📊 Ordenar por:",
            options=["PROMEDIO_PONDERADO_GENERAL"] + metricas_materias,
            format_func=lambda x: "Promedio Ponderado General" if x == "PROMEDIO_PONDERADO_GENERAL" else x.replace("_PROMEDIO_GENERAL", "").replace("_", " "),
            index=0,
        )
    with col2:
        min_puntaje = st.number_input("📉 Puntaje mínimo:", min_value=0.0, max_value=500.0, value=0.0, step=10.0)
    with col3:
        max_puntaje = st.number_input("📈 Puntaje máximo:", min_value=0.0, max_value=500.0, value=500.0, step=10.0)
    with col4:
        grados_disponibles = sorted(set(grados_list)) if grados_list else []
        grado_filtro = st.multiselect("🎓 Filtrar por grado:", options=["Todos"] + grados_disponibles, default=["Todos"])

    col1, col2 = st.columns(2)
    with col1:
        simulacros_mostrar = st.multiselect(
            "📋 Mostrar columnas de simulacros:",
            options=[sim["nombre"] for sim in simulacros],
            default=[sim["nombre"] for sim in simulacros],
        )
    with col2:
        buscar_nombre = st.text_input("🔍 Buscar estudiante por nombre:", placeholder="Escribe el nombre...")

    tabla_filtrada = datos_unificados.copy()
    tabla_filtrada = tabla_filtrada[(tabla_filtrada[metrica_ordenar] >= min_puntaje) & (tabla_filtrada[metrica_ordenar] <= max_puntaje)]
    if "GRADO" in tabla_filtrada.columns and "Todos" not in grado_filtro:
        tabla_filtrada = tabla_filtrada[tabla_filtrada["GRADO"].astype(str).isin(grado_filtro)]
    if buscar_nombre:
        tabla_filtrada = tabla_filtrada[tabla_filtrada["ESTUDIANTE"].str.contains(buscar_nombre.upper(), na=False)]

    tabla_filtrada = tabla_filtrada.sort_values(metrica_ordenar, ascending=False).reset_index(drop=True)
    tabla_filtrada.insert(0, "RANKING", range(1, len(tabla_filtrada) + 1))

    columnas_mostrar = ["RANKING", "ESTUDIANTE"]
    if "GRADO" in tabla_filtrada.columns:
        columnas_mostrar.append("GRADO")
    columnas_mostrar.append("PROMEDIO_PONDERADO_GENERAL")

    # Añadir columnas de simulacros seleccionados
    for sim in simulacros:
        slug = _slug(sim["nombre"])
        if sim["nombre"] in simulacros_mostrar:
            prom_col = f"{slug}_PROM"
            if prom_col in tabla_filtrada.columns:
                columnas_mostrar.append(prom_col)
            for mat in materias:
                col_mat = f"{mat}_{slug}"
                if col_mat in tabla_filtrada.columns:
                    columnas_mostrar.append(col_mat)

    columnas_mostrar.extend(metricas_materias)
    columnas_mostrar = [c for c in columnas_mostrar if c in tabla_filtrada.columns]
    tabla_mostrar = tabla_filtrada[columnas_mostrar].copy()
    columnas_numericas = tabla_mostrar.select_dtypes(include=[np.number]).columns
    tabla_mostrar[columnas_numericas] = tabla_mostrar[columnas_numericas].round(2)

    st.markdown("### 📊 Estadísticas del Ranking Filtrado")
    col1, col2, col3, col4 = st.columns(4)
    with col1:
        st.metric("👥 Total Estudiantes", len(tabla_filtrada))
    with col2:
        st.metric("📈 Promedio", f"{tabla_filtrada[metrica_ordenar].mean():.2f}")
    with col3:
        st.metric("🏆 Máximo", f"{tabla_filtrada[metrica_ordenar].max():.2f}")
    with col4:
        st.metric("📉 Mínimo", f"{tabla_filtrada[metrica_ordenar].min():.2f}")

    st.markdown("---")
    tab1, tab2, tab3 = st.tabs(["📋 Tabla Completa", "📊 Gráfica de Ranking", "📈 Distribución"])

    with tab1:
        st.markdown("#### 📋 Ranking Completo")
        tabla_display = tabla_mostrar.copy()
        rename_dict = {}
        for col in tabla_display.columns:
            if col.endswith("_PROM") and col != "PROMEDIO_PONDERADO_GENERAL":
                name_clean = col.replace("_PROM", "").replace("_", " ").title()
                rename_dict[col] = f"PROM. {name_clean}"
            elif "_PROMEDIO_GENERAL" in col:
                rename_dict[col] = col.replace("_PROMEDIO_GENERAL", " (Promedio)").replace("_", " ")
            elif col == "PROMEDIO_PONDERADO_GENERAL":
                rename_dict[col] = "PROMEDIO GENERAL"
        tabla_display = tabla_display.rename(columns=rename_dict)
        columnas_para_gradiente = [col for col in tabla_display.columns if col not in ["RANKING", "ESTUDIANTE", "GRADO"]]

        def resaltar_promedio(col):
            if col.name == "PROMEDIO GENERAL":
                return ["background-color: #007BFF; color: white; font-weight: bold"] * len(col)
            return [""] * len(col)

        tabla_estilo = (
            tabla_display.style.background_gradient(subset=columnas_para_gradiente, cmap="RdYlGn", vmin=0, vmax=100)
            .apply(resaltar_promedio)
            .format({col: "{:.2f}" for col in columnas_para_gradiente})
        )
        st.dataframe(tabla_estilo, width="stretch", height=600)

    with tab2:
        st.markdown("#### 📊 Top 30 Estudiantes")
        top_30 = tabla_filtrada.head(30)
        custom_colorscale = [
            [0.0, "#E74C3C"],
            [0.4, "#D35400"],
            [0.68, "#3498DB"],
            [0.8, "#2ECC71"],
            [1.0, "#27AE60"],
        ]
        fig = go.Figure()
        fig.add_trace(
            go.Bar(
                y=top_30["ESTUDIANTE"],
                x=top_30[metrica_ordenar],
                orientation="h",
                marker=dict(color=top_30[metrica_ordenar], colorscale=custom_colorscale, cmin=0, cmax=500, showscale=True, colorbar=dict(title="Puntaje")),
                text=top_30[metrica_ordenar].round(2),
                textposition="outside",
            )
        )
        fig.update_layout(
            title=f"Top 30 - {metrica_ordenar.replace('_', ' ')}",
            xaxis_title="Puntaje",
            yaxis_title="Estudiante",
            height=800,
            yaxis={"categoryorder": "total ascending"},
            template="plotly_white",
        )
        st.plotly_chart(fig, width="stretch")

    with tab3:
        st.markdown("#### 📈 Distribución de Puntajes")
        col1, col2 = st.columns(2)
        with col1:
            fig = go.Figure()
            fig.add_trace(go.Histogram(x=tabla_filtrada[metrica_ordenar], nbinsx=30, marker_color="#667eea", name="Frecuencia"))
            fig.add_vline(
                x=tabla_filtrada[metrica_ordenar].mean(),
                line_dash="dash",
                line_color="red",
                annotation_text=f"Promedio: {tabla_filtrada[metrica_ordenar].mean():.2f}",
            )
            fig.update_layout(title="Distribución de Puntajes", xaxis_title="Puntaje", yaxis_title="Frecuencia", height=400, template="plotly_white")
            st.plotly_chart(fig, width="stretch")
        with col2:
            fig = go.Figure()
            fig.add_trace(go.Box(y=tabla_filtrada[metrica_ordenar], name="Puntajes", marker_color="#667eea", boxmean="sd"))
            fig.update_layout(title="Estadísticas de Puntajes", yaxis_title="Puntaje", height=400, template="plotly_white")
            st.plotly_chart(fig, width="stretch")

    st.markdown("---")
    st.markdown("### 💾 Descargar Datos")

    def num_to_excel_col(n):
        result = ""
        while n > 0:
            n -= 1
            result = chr(65 + (n % 26)) + result
            n //= 26
        return result

    col1, col2 = st.columns(2)
    with col1:
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.utils.dataframe import dataframe_to_rows
        from openpyxl.worksheet.table import Table, TableStyleInfo

        output = BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Ranking Filtrado"

        for r in dataframe_to_rows(tabla_mostrar, index=False, header=True):
            ws.append(r)

        last_col = num_to_excel_col(len(tabla_mostrar.columns))
        last_row = len(tabla_mostrar) + 1
        tab = Table(displayName="TablaRanking", ref=f"A1:{last_col}{last_row}")
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)

        wb.save(output)
        excel_data = output.getvalue()

        st.download_button(
            label="📥 Descargar Tabla Filtrada (Excel)",
            data=excel_data,
            file_name="ranking_filtrado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    with col2:
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.utils.dataframe import dataframe_to_rows
        from openpyxl.worksheet.table import Table, TableStyleInfo

        output_completo = BytesIO()
        wb_completo = Workbook()
        ws_completo = wb_completo.active
        ws_completo.title = "Datos Completos"

        for r in dataframe_to_rows(datos_unificados, index=False, header=True):
            ws_completo.append(r)

        last_col_completo = num_to_excel_col(len(datos_unificados.columns))
        last_row_completo = len(datos_unificados) + 1
        tab_completo = Table(displayName="DatosCompletos", ref=f"A1:{last_col_completo}{last_row_completo}")
        style_completo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tab_completo.tableStyleInfo = style_completo
        ws_completo.add_table(tab_completo)

        wb_completo.save(output_completo)
        excel_completo = output_completo.getvalue()

        st.download_button(
            label="📥 Descargar Datos Completos (Excel)",
            data=excel_completo,
            file_name="datos_completos_todos_simulacros.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
